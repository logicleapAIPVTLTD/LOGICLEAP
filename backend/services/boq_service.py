import google.generativeai as genai
import json
import logging
import pdfplumber
from docx import Document
import io

logger = logging.getLogger(__name__)

class BOQService:
    def __init__(self, api_key: str, model_name: str = "gemini-2.5-flash-lite"):
        genai.configure(api_key=api_key)
        self.model = genai.GenerativeModel(model_name)
        self.generation_config = genai.types.GenerationConfig(
            temperature=0.1,
            max_output_tokens=8192,
        )

    def extract_text(self, file_bytes: bytes, filename: str) -> str:
        try:
            ext = filename.split('.')[-1].lower()
            if ext == "pdf":
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    return "\n".join([p.extract_text() or "" for p in pdf.pages])
            elif ext == "docx":
                doc = Document(io.BytesIO(file_bytes))
                return "\n".join([p.text for p in doc.paragraphs])
            elif ext == "txt":
                return file_bytes.decode("utf-8")
        except Exception as e:
            logger.error(f"File extraction failed: {e}")
            return ""
        return ""

    def get_identification_prompt(self, context: dict) -> str:
        p_type = context.get('project_type', 'Interior')
        location = context.get('location', 'General')
        
        return f"""
        Role: Senior Quantity Surveyor & Interior Estimator.
        
        PROJECT CONTEXT:
        - Project: {context.get('project_name')}
        - Category: {p_type}
        - Location: {location}

        PRE-TASK (IF NEEDED): If the input image/floor plan is too big or complex try thinking more and break it down into smaller sections for better analysis upscale and think to give the best possible output but follow the strict output format. 

        TASK:
        1. COMPREHENSIVE IDENTIFICATION: For every room identified, generate a FULL PACKAGE of work.
        2. STANDARD ITEMS TO INCLUDE (IF INTERIOR):
           - Flooring (PU/Tiles/Marble)
           - Skirting (matched to flooring)
           - False Ceiling (Gypsum/Grid)
           - Internal Wall Painting (Plastic Emulsion/Luster)
           - Electrical Points (Switchboards/Wiring placeholders)
           - Woodwork/Carpentry (Wardrobes, TV Units, or Paneling)
           - Doors & Windows (Frames and Shutters)
        
        3. DIMENSIONS: Extract or infer Length and Width. If only area is given, derive L & W (e.g., 100 sqft -> 10x10).
        4. TIER: Assign Tier (T1/T2/T3) based on city profile.
        
        STRICT DOMAIN GUARDRAIL:
        - Ignore all road, highway, or external infrastructure data. 
        - If input is an empty floor plan, use standard architectural sizes for the {p_type} project.

        STRICT OUTPUT SCHEMA (JSON LIST ONLY):
        [
          {{
            "Item No.": integer,
            "Work": "string (Room Name + Specific Work, e.g., Master Bedroom Plastic Emulsion Painting)",
            "State": "string",
            "Tier": "string",
            "Length": number,
            "Width": number,
            "Quantity": number,
            "Unit": "string"
          }}
        ]
        """

    def process(self, content: str, context: dict, image_parts=None):
        sys_prompt = self.get_identification_prompt(context)
        
        try:
            logger.info("Sending request to Gemini...")
            if image_parts:
                response = self.model.generate_content([image_parts[0], sys_prompt], generation_config=self.generation_config)
            else:
                response = self.model.generate_content(f"{sys_prompt}\n\nINPUT DATA:\n{content}", generation_config=self.generation_config)

            raw_text = response.text.replace("```json", "").replace("```", "").strip()
            
            start = raw_text.find("[")
            end = raw_text.rfind("]") + 1
            
            if start == -1 or end == 0:
                logger.error(f"Invalid JSON received from AI: {raw_text[:100]}...")
                return []
            
            json_str = raw_text[start:end]
            return json.loads(json_str)
        except Exception as e:
            logger.error(f"❌ Identification Error: {e}")
            return []